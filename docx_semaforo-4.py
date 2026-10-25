"""Genera el Semáforo (Evaluación Rápida de Daños Post-Sísmico) en Word,
sobre la hoja membretada de la ficha (assets/membrete_semaforo.docx).

El membrete (logos, franjas, marca de agua, pie y las tres direcciones:
Secretaría, Dirección de Ingeniería Municipal y Dirección de Hábitat y Vivienda)
vive en el ENCABEZADO de la plantilla, así que sale en todas las páginas.
Usa los mismos datos, textos y lógica de resolución de logic_semaforo.py.
"""
import os
from itertools import zip_longest

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import logic_semaforo as ls

RUTA_PLANTILLA = os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets", "membrete_semaforo.docx")

# Paleta y medidas tomadas de la ficha técnica
AZUL = "1A365D"
GRIS_ETQ = "E2E8F0"
BORDE_SUAVE = "CBD5E0"
ROJO, AMARILLO, VERDE = "C0392B", "D68910", "27AE60"
ANCHO_TABLA = 10600  # dxa (área útil entre márgenes de 800)
FUENTE = "Arial"


# ---------------------------------------------------------------- utilidades XML
def _fuente(run, size=8.5, bold=False, italic=False, color=None, underline=False):
    run.font.name = FUENTE
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for att in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(att), FUENTE)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.underline = underline
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _parrafo(par, antes=0, despues=0, alinea=None, interlineado=None):
    pf = par.paragraph_format
    pf.space_before = Pt(antes)
    pf.space_after = Pt(despues)
    if interlineado:
        pf.line_spacing = Pt(interlineado)
    if alinea is not None:
        par.alignment = alinea


def _sombrear(celda, fill):
    tcpr = celda._tc.get_or_add_tcPr()
    for viejo in tcpr.findall(qn("w:shd")):
        tcpr.remove(viejo)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def _bordes_tabla(tabla, externo, interno, sz_ext=6, sz_int=4, interior=True):
    tblpr = tabla._tbl.tblPr
    for viejo in tblpr.findall(qn("w:tblBorders")):
        tblpr.remove(viejo)
    b = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz_ext)); e.set(qn("w:space"), "0"); e.set(qn("w:color"), externo)
        b.append(e)
    for lado in ("insideH", "insideV"):
        e = OxmlElement(f"w:{lado}")
        if interior:
            e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz_int)); e.set(qn("w:space"), "0"); e.set(qn("w:color"), interno)
        else:
            e.set(qn("w:val"), "nil")
        b.append(e)
    tblpr.append(b)


def _formato_tabla(tabla, anchos, margen_h=120, margen_v=60):
    """Ancho fijo, centrada, anchos en dxa en la cuadrícula y en cada celda."""
    tblpr = tabla._tbl.tblPr
    tabla.autofit = False
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    for viejo in tblpr.findall(qn("w:tblW")):
        tblpr.remove(viejo)
    w = OxmlElement("w:tblW"); w.set(qn("w:type"), "dxa"); w.set(qn("w:w"), str(sum(anchos)))
    tblpr.append(w)
    lay = tblpr.find(qn("w:tblLayout"))
    if lay is None:
        lay = OxmlElement("w:tblLayout"); tblpr.append(lay)
    lay.set(qn("w:type"), "fixed")
    mar = OxmlElement("w:tblCellMar")
    for lado, val in (("top", margen_v), ("left", margen_h), ("bottom", margen_v), ("right", margen_h)):
        e = OxmlElement(f"w:{lado}"); e.set(qn("w:w"), str(val)); e.set(qn("w:type"), "dxa"); mar.append(e)
    tblpr.append(mar)
    grid = tabla._tbl.tblGrid
    for gc, ancho in zip(grid.findall(qn("w:gridCol")), anchos):
        gc.set(qn("w:w"), str(ancho))
    for fila in tabla.rows:
        trpr = fila._tr.get_or_add_trPr()
        if trpr.find(qn("w:cantSplit")) is None:
            trpr.append(OxmlElement("w:cantSplit"))
        for celda, ancho in zip(fila.cells, anchos):
            tcpr = celda._tc.get_or_add_tcPr()
            for viejo in tcpr.findall(qn("w:tcW")):
                tcpr.remove(viejo)
            tcw = OxmlElement("w:tcW"); tcw.set(qn("w:type"), "dxa"); tcw.set(qn("w:w"), str(ancho))
            tcpr.insert(0, tcw)


def _texto_celda(celda, texto, size=8.5, bold=False, color=None, italic=False, alinea=None, valign=None):
    par = celda.paragraphs[0]
    _parrafo(par, alinea=alinea)
    _fuente(par.add_run(texto), size=size, bold=bold, color=color, italic=italic)
    if valign is not None:
        celda.vertical_alignment = valign
    return par


def _hipervinculo(par, url, texto, size=8.5):
    r_id = par.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hl = OxmlElement("w:hyperlink")
    hl.set(qn("r:id"), r_id)
    hl.set(qn("w:history"), "1")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rf = OxmlElement("w:rFonts")
    for att in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(att), FUENTE)
    rpr.append(rf)
    col = OxmlElement("w:color"); col.set(qn("w:val"), "1155CC"); rpr.append(col)
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(int(size * 2))); rpr.append(sz)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
    run.append(rpr)
    t = OxmlElement("w:t"); t.text = texto; t.set(qn("xml:space"), "preserve")
    run.append(t)
    hl.append(run)
    par._p.append(hl)


def _titulo_seccion(doc, texto, salto_pagina=False):
    p = doc.add_paragraph()
    _parrafo(p, antes=8, despues=4)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = salto_pagina
    _fuente(p.add_run(texto), size=10, bold=True, color=AZUL)
    return p


def _espacio(doc, pts=4):
    p = doc.add_paragraph()
    _parrafo(p, interlineado=pts)
    return p


# ---------------------------------------------------------------- secciones
def _tabla_datos_generales(doc, datos):
    g = datos["generales"]
    adultos = int(datos.get("adultos") or 0)
    menores = int(datos.get("menores") or 0)
    lat = (datos.get("latitud") or "").strip()
    lon = (datos.get("longitud") or "").strip()
    inspector = (datos.get("inspector") or "").strip()

    gps = ("link", f"https://www.google.com/maps?q={lat},{lon}", f"Ver ubicación satelital ({lat}, {lon})") if lat and lon else "No registrada"

    # Columna izquierda = ocupante / vivienda; derecha = construcción y estado.
    # Con colapso se omiten las condiciones que ya no se evalúan.
    colapsado = ls.colapso_activo(g)
    izq = [
        ("PARROQUIA:", (datos.get("parroquia", "") + ", Brión")),
        ("OCUPANTE:", datos.get("ocupante") or "Sin Identificar"),
        ("TELÉFONO:", datos.get("telefono") or "No registrado"),
        ("CÉDULA DEL BENEFICIARIO:", datos.get("cedula") or "No registrada"),
        ("PERSONAS QUE HABITAN:", f"{adultos + menores} ({adultos} adultos, {menores} menores)"),
        ("NÚMERO DE NIVELES:", datos.get("niveles") or "No registrado"),
        ("OCUPACIÓN ACTUAL:", datos.get("ocupacion", "")),
        ("TENENCIA:", datos.get("tenencia") or "No registrada"),
        ("FECHA:", datos.get("fecha", "")),
    ]
    if inspector:
        izq.append(("INSPECTOR:", inspector))
    der = [
        ("TECNOLOGÍA CONSTRUCTIVA:", datos.get("tecnologia_constructiva", "")),
        ("TIPOLOGÍA:", datos.get("tipologia", "")),
        ("UBICACIÓN GPS:", gps),
        ("COLAPSO DE LA ESTRUCTURA:", g.get("colapso") or ls.COLAPSO_NINGUNO),
    ]
    if not colapsado:
        der.extend([
            ("ASENTAMIENTO DEL TERRENO:", g["asentamiento"]),
            ("INCLINACIÓN / DESPLOME:", g["inclinacion"]),
            ("ESTRUCTURA VECINA EN RIESGO:", g["vecino_riesgo"]),
            ("RIESGO DE CAÍDA (cables, tanques, muros):", g["riesgo_caida"]),
            ("SERVICIOS BÁSICOS:", g["servicios"]),
        ])
    filas = [[("DIRECCIÓN:", datos.get("vivienda", ""))]]
    filas += [[a, b] for a, b in zip_longest(izq, der)]

    anchos = [2300, 3000, 2300, 3000]
    tabla = doc.add_table(rows=len(filas), cols=4)
    for i, fila in enumerate(filas):
        celdas = tabla.rows[i].cells
        if len(fila) == 1:
            etq, val = fila[0]
            fusion = celdas[1].merge(celdas[3])
            _sombrear(celdas[0], GRIS_ETQ)
            _texto_celda(celdas[0], etq, bold=True, color=AZUL)
            _texto_celda(fusion, str(val))
            continue
        for j, item in enumerate(fila):
            ce, cv = celdas[j * 2], celdas[j * 2 + 1]
            if item is None:  # celda vacía, sin sombreado
                continue
            etq, val = item
            _sombrear(ce, GRIS_ETQ)
            _texto_celda(ce, etq, bold=True, color=AZUL)
            if isinstance(val, tuple) and val[0] == "link":
                par = cv.paragraphs[0]
                _parrafo(par)
                _hipervinculo(par, val[1], val[2])
            else:
                _texto_celda(cv, str(val))
    _formato_tabla(tabla, anchos)
    _bordes_tabla(tabla, AZUL, BORDE_SUAVE)
    return tabla


def _tabla_espacios(doc, datos):
    """Devuelve la tabla y la lista de (nombre, ruta_foto) para el registro fotográfico.
    Recorre ls.ESPACIOS_REGULARES, que incluye las cuatro fachadas."""
    fotos = []
    filas_txt = []
    for clave, nombre in ls.ESPACIOS_REGULARES:
        info = datos["espacios"].get(clave, {"estructural": [], "no_estructural": [], "cantidades_estructural": {}, "cantidades_no_estructural": {}, "nota": "", "foto": None})
        cant_e = info.get("cantidades_estructural", {})
        cant_n = info.get("cantidades_no_estructural", {})
        sev = max([ls._RANGO_ESTRUCTURAL.get(v, 0) for v in info["estructural"]], default=0)

        def _con_cant(valores, mapa, cantidades):
            partes = []
            for v in valores:
                t = mapa[v]
                c = cantidades.get(v)
                if c and str(c).strip() not in ("", "1"):
                    t += f" (x{c})"
                partes.append(t)
            return ", ".join(partes)

        t_e = _con_cant(info["estructural"], ls._MAPA_ESTRUCTURAL, cant_e) or "Sin daño"
        t_n = _con_cant(info["no_estructural"], ls._MAPA_NO_ESTRUCTURAL, cant_n) or "Sin daño"
        c_e = ROJO if sev >= 2 else (AMARILLO if sev == 1 else VERDE)
        c_n = ROJO if "colapso" in info["no_estructural"] else (AMARILLO if info["no_estructural"] else VERDE)
        filas_txt.append((nombre, (t_e, c_e), (t_n, c_n), info.get("nota") or "-"))
        if info.get("foto") and os.path.isfile(info["foto"]):
            fotos.append((nombre, info["foto"]))

    techo = datos.get("techo", {"marcado": [], "nota": "", "foto": None})
    t_t = ", ".join(ls._MAPA_TECHO[v] for v in techo["marcado"]) if techo["marcado"] else "Sin daño"
    c_t = ROJO if ("colapso" in techo["marcado"] or "parcial" in techo["marcado"]) else (AMARILLO if techo["marcado"] else VERDE)
    filas_txt.append(("TECHO / CUBIERTA", ("—", None), (t_t, c_t), techo.get("nota") or "-"))
    if techo.get("foto") and os.path.isfile(techo["foto"]):
        fotos.append(("Techo / Cubierta", techo["foto"]))

    anchos = [1900, 3000, 3000, 2700]
    tabla = doc.add_table(rows=1 + len(filas_txt), cols=4)
    for j, h in enumerate(["Espacio", "Daño estructural", "Daño no estructural", "Nota / ubicación"]):
        c = tabla.rows[0].cells[j]
        _sombrear(c, AZUL)
        _texto_celda(c, h, bold=True, color="FFFFFF", alinea=WD_ALIGN_PARAGRAPH.CENTER, valign=WD_ALIGN_VERTICAL.CENTER)
    for i, (nombre, (te, ce), (tn, cn), nota) in enumerate(filas_txt, start=1):
        fila = tabla.rows[i].cells
        _texto_celda(fila[0], nombre, valign=WD_ALIGN_VERTICAL.CENTER)
        _texto_celda(fila[1], te, bold=ce is not None, color=ce, valign=WD_ALIGN_VERTICAL.CENTER)
        _texto_celda(fila[2], tn, bold=True, color=cn, valign=WD_ALIGN_VERTICAL.CENTER)
        _texto_celda(fila[3], nota, valign=WD_ALIGN_VERTICAL.CENTER)
    _formato_tabla(tabla, anchos, margen_v=80)
    _bordes_tabla(tabla, AZUL, BORDE_SUAVE)
    return tabla, fotos


def _caja_resolucion(doc, color_hex, etiqueta, motivo):
    color = color_hex.lstrip("#")
    tabla = doc.add_table(rows=2, cols=1)
    _sombrear(tabla.rows[0].cells[0], color)
    _texto_celda(tabla.rows[0].cells[0], f"RESOLUCIÓN: {etiqueta}", size=11, bold=True, color="FFFFFF")
    par = tabla.rows[1].cells[0].paragraphs[0]
    _parrafo(par)
    _fuente(par.add_run("Motivo: "), size=9, bold=True)
    _fuente(par.add_run(motivo), size=9)
    _formato_tabla(tabla, [ANCHO_TABLA], margen_h=160, margen_v=90)
    _bordes_tabla(tabla, color, color, sz_ext=16, interior=False)
    for fila in tabla.rows:  # que la caja no se parta entre páginas
        for c in fila.cells:
            for p in c.paragraphs:
                p.paragraph_format.keep_with_next = True


def _galeria(doc, fotos, lado_cm=8.0):
    _titulo_seccion(doc, "III. REGISTRO FOTOGRÁFICO", salto_pagina=True)
    filas = [fotos[i:i + 2] for i in range(0, len(fotos), 2)]
    tabla = doc.add_table(rows=len(filas), cols=2)
    for i, grupo in enumerate(filas):
        for j in range(2):
            celda = tabla.rows[i].cells[j]
            celda.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            if j >= len(grupo):
                _texto_celda(celda, "")
                continue
            nombre, ruta = grupo[j]
            par = celda.paragraphs[0]
            _parrafo(par, antes=4, despues=2, alinea=WD_ALIGN_PARAGRAPH.CENTER)
            try:
                par.add_run().add_picture(ls.recortar_cuadrado(ruta), width=Cm(lado_cm), height=Cm(lado_cm))
            except Exception:
                _fuente(par.add_run("(no se pudo cargar la imagen)"), size=8.5, italic=True)
            cap = celda.add_paragraph()
            _parrafo(cap, despues=6, alinea=WD_ALIGN_PARAGRAPH.CENTER)
            _fuente(cap.add_run(nombre), size=9, bold=True, color=AZUL)
    _formato_tabla(tabla, [5300, 5300])
    _bordes_tabla(tabla, "FFFFFF", "FFFFFF", sz_ext=0, sz_int=0, interior=False)
    # sin bordes visibles
    tblpr = tabla._tbl.tblPr
    for b in tblpr.findall(qn("w:tblBorders")):
        for e in list(b):
            e.set(qn("w:val"), "nil")



_ORDEN_TBLPR = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize",
                "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook",
                "tblCaption", "tblDescription"]
_ORDEN_TCPR = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
               "textDirection", "tcFitText", "vAlign", "hideMark"]


def _ordenar_hijos(elemento, orden):
    """Reordena los hijos según el esquema OOXML (Word es estricto con el orden)."""
    def clave(hijo):
        nombre = hijo.tag.split("}")[1]
        return orden.index(nombre) if nombre in orden else len(orden)
    hijos = sorted(list(elemento), key=clave)
    for h in list(elemento):
        elemento.remove(h)
    for h in hijos:
        elemento.append(h)


def _normalizar_tablas(doc):
    for tbl in doc.element.body.iter(qn("w:tbl")):
        tblpr = tbl.find(qn("w:tblPr"))
        if tblpr is not None:
            _ordenar_hijos(tblpr, _ORDEN_TBLPR)
        for tcpr in tbl.iter(qn("w:tcPr")):
            _ordenar_hijos(tcpr, _ORDEN_TCPR)

# ---------------------------------------------------------------- API pública
def ruta_docx(identificador, carpeta):
    return ls.ruta_pdf(identificador, carpeta)[:-4] + ".docx"


def generar_docx(datos, carpeta):
    """Genera el Word del semáforo. Devuelve (ruta_docx, color_hex, etiqueta)."""
    if not os.path.isfile(RUTA_PLANTILLA):
        raise FileNotFoundError(f"Falta la plantilla membretada: {RUTA_PLANTILLA}")
    doc = Document(RUTA_PLANTILLA)

    t = doc.add_paragraph()
    _parrafo(t, despues=8, alinea=WD_ALIGN_PARAGRAPH.CENTER)
    _fuente(t.add_run("EVALUACIÓN RÁPIDA DE DAÑOS POST-SÍSMICO (SEMÁFORO)"), size=10.5, bold=True, color=AZUL)

    _titulo_seccion(doc, "I. DATOS GENERALES")
    _tabla_datos_generales(doc, datos)

    _titulo_seccion(doc, "II. EVALUACIÓN POR ESPACIO")
    if ls.colapso_activo(datos["generales"]):
        fotos = []
        p = doc.add_paragraph()
        _parrafo(p, despues=4)
        _fuente(p.add_run("No se realizó la evaluación por espacio porque se reportó colapso de la estructura."), size=8.5, italic=True)
    else:
        _, fotos = _tabla_espacios(doc, datos)
    if datos.get("foto_colapso") and os.path.isfile(datos["foto_colapso"]):
        fotos.insert(0, ("Colapso de la estructura", datos["foto_colapso"]))

    _espacio(doc, 8)
    techo = datos.get("techo", {"marcado": []})
    color_hex, etiqueta, motivo = ls.calcular_semaforo(datos["espacios"], techo["marcado"], datos["generales"])
    _caja_resolucion(doc, color_hex, etiqueta, motivo)

    nota = doc.add_paragraph()
    _parrafo(nota, antes=4)
    _fuente(nota.add_run("Esta clasificación es una evaluación visual preliminar. La decisión de demolición corresponde exclusivamente a una evaluación estructural formal."), size=7.5, italic=True, color="5D6D7E")

    inspeccion, medidas = ls.acciones_recomendadas(datos)
    if inspeccion or medidas:
        _titulo_seccion(doc, "ACCIONES RECOMENDADAS")
        for titulo, items in (("Inspección detallada: ", inspeccion), ("Medidas de prevención: ", medidas)):
            if items:
                p = doc.add_paragraph()
                _parrafo(p, despues=2)
                _fuente(p.add_run(titulo), size=8.5, bold=True)
                _fuente(p.add_run(", ".join(items)), size=8.5)

    obs = (datos.get("observaciones_generales") or "").strip()
    if obs:
        _titulo_seccion(doc, "OBSERVACIONES GENERALES")
        for linea in obs.splitlines():
            p = doc.add_paragraph()
            _parrafo(p, despues=2)
            _fuente(p.add_run(linea), size=8.5)

    if fotos:
        _galeria(doc, fotos)

    _normalizar_tablas(doc)
    salida = ruta_docx(datos.get("vivienda", "SIN_NOMBRE"), carpeta)
    doc.save(salida)
    return salida, color_hex, etiqueta
